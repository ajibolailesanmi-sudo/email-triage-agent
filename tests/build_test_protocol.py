"""Build tests/test-protocol.xlsx from the recorded results.

Inputs: eval/eval-set.csv (the labelled emails), eval/results-prompt-1.1.csv (what the 02 scenario wrote for
prompt 1.1) and tests/test_protocol_data.py (live checks, manual tests, acceptance criteria).

    python3 tests/build_test_protocol.py [output.xlsx]
"""
import csv
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from test_protocol_data import (  # noqa: E402
    ACCEPTANCE, CHANGES, DETAILS, EARLIER_RUNS, EV_CHECKS, FIXED, HOW_TO, LIVE, LIVE_EXAMPLE, MANUAL,
    MANUAL_EXAMPLE, MANUAL_NOTE, OPEN_ITEMS, READINGS, RUN_11_OPERATIONS,
)

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "test-protocol.xlsx")

F = "Arial"
TITLE = Font(name=F, size=16, bold=True, color="1F3A5F")
SUB = Font(name=F, size=10, italic=True, color="666666")
SECTION = Font(name=F, size=11, bold=True, color="1F3A5F")
BOLD = Font(name=F, size=10, bold=True)
BODY = Font(name=F, size=10)
HEAD = Font(name=F, size=10, bold=True, color="FFFFFF")
EXAMPLE = Font(name=F, size=10, italic=True, color="808080")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
EX_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="D9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top")
GREEN = (PatternFill("solid", fgColor="C6EFCE"), Font(name=F, size=10, color="006100", bold=True))
RED = (PatternFill("solid", fgColor="FFC7CE"), Font(name=F, size=10, color="9C0006", bold=True))
AMBER = (PatternFill("solid", fgColor="FFEB9C"), Font(name=F, size=10, color="9C5700", bold=True))


def style(ws, rng, font=BODY, fill=None, border=True, align=WRAP, fmt=None):
    for row in ws[rng]:
        for c in row:
            c.font = font
            if fill:
                c.fill = fill
            if border:
                c.border = BOX
            if align:
                c.alignment = align
            if fmt:
                c.number_format = fmt


def header(ws, row, labels, col=1):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.font, c.fill, c.border = HEAD, HEAD_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def fit_row(ws, row, first_col=1, last_col=None, merged_width=None, minimum=15):
    """Set a row height that shows all wrapped text (Excel and Sheets do not always autofit)."""
    last_col = last_col or ws.max_column
    lines = 1
    for col in range(first_col, last_col + 1):
        v = ws.cell(row=row, column=col).value
        if v is None or (isinstance(v, str) and v.startswith("=")):
            continue
        width = merged_width or ws.column_dimensions[get_column_letter(col)].width or 9
        per_line = max(1, int(width * 1.05))
        count = sum(max(1, math.ceil(len(part) / per_line)) for part in str(v).split("\n"))
        lines = max(lines, count)
    ws.row_dimensions[row].height = max(minimum, 13 * lines + 3)


def pass_fail(ws, rng, validate=None):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], fill=GREEN[0], font=GREEN[1]))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], fill=RED[0], font=RED[1]))
    if validate:
        dv = DataValidation(type="list", formula1='"PASS,FAIL"', allow_blank=True, promptTitle="Result",
                            prompt="PASS or FAIL. Leave blank if not run yet.", showInputMessage=True)
        ws.add_data_validation(dv)
        dv.add(validate)


def met_rules(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Met"'], fill=GREEN[0], font=GREEN[1]))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Not met"'], fill=RED[0], font=RED[1]))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Pending"'], fill=AMBER[0], font=AMBER[1]))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Partly met"'], fill=AMBER[0], font=AMBER[1]))


def landscape(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def title(ws, text, sub, merge_to):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A2"] = sub
    ws["A2"].font = SUB
    ws["A2"].alignment = WRAP
    ws.merge_cells(f"A2:{merge_to}2")


def as_bool(v):
    return {"TRUE": True, "FALSE": False}.get(v.strip().upper(), v)


def as_number(v):
    try:
        return int(v)
    except ValueError:
        return float(v)


emails = {r["id"]: r for r in csv.DictReader(open(os.path.join(ROOT, "eval", "eval-set.csv"), encoding="utf-8"))}
results = list(csv.DictReader(open(os.path.join(ROOT, "eval", "results-prompt-1.1.csv"), encoding="utf-8")))

wb = Workbook()

# ---------------- Live checks ----------------
lv = wb.active
lv.title = "Live checks"
title(lv, "Live integration checks",
      "03 - Seed Test Emails sends labelled emails to the inbox and 01 - Triage Inbox processes them as real mail; "
      "two checks (L-05, L-06) happened for real. Yellow columns hold the recorded result. Row 5 is an example.",
      "I")
lv.row_dimensions[2].height = 30
header(lv, 4, ["Test ID", "Check", "Set-up", "Expected result", "Acceptance criteria", "Actual result",
               "Pass/Fail", "Evidence", "Date"])
widths(lv, [9, 24, 40, 46, 13, 50, 10, 32, 11])
lv.append(list(LIVE_EXAMPLE))
style(lv, "A5:I5", font=EXAMPLE, fill=EX_FILL)
for row in LIVE:
    lv.append(list(row))
last_l = 5 + len(LIVE)
style(lv, f"A6:E{last_l}")
style(lv, f"F6:I{last_l}", fill=INPUT_FILL)
for r in range(5, last_l + 1):
    lv.cell(row=r, column=1).font = BOLD if r > 5 else EXAMPLE
    lv.cell(row=r, column=7).alignment = CENTER
    fit_row(lv, r)
pass_fail(lv, f"G5:G{last_l}", validate=f"G6:G{last_l}")
lv.freeze_panes = "B5"
landscape(lv)

# ---------------- Manual ----------------
mn = wb.create_sheet("Manual")
title(mn, "Manual and failure tests",
      "Tests a person runs or a setting is changed for. Yellow columns hold the recorded result. Row 5 is an example.",
      "J")
header(mn, 4, ["Test ID", "Scenario", "Steps", "Expected result", "Acceptance criteria", "Actual result",
               "Pass/Fail", "Evidence", "Tested by", "Date"])
widths(mn, [9, 24, 40, 42, 13, 50, 10, 26, 22, 11])
mn.append(list(MANUAL_EXAMPLE))
style(mn, "A5:J5", font=EXAMPLE, fill=EX_FILL)
for row in MANUAL:
    mn.append(list(row))
last_m = 5 + len(MANUAL)
style(mn, f"A6:E{last_m}")
style(mn, f"F6:J{last_m}", fill=INPUT_FILL)
for r in range(5, last_m + 1):
    mn.cell(row=r, column=1).font = BOLD if r > 5 else EXAMPLE
    mn.cell(row=r, column=7).alignment = CENTER
    fit_row(mn, r)
pass_fail(mn, f"G5:G{last_m}", validate=f"G6:G{last_m}")
mn.freeze_panes = "B5"
landscape(mn)

# ---------------- Evaluation ----------------
ev = wb.create_sheet("Evaluation")
title(ev, "Classifier evaluation, prompt 1.1",
      "The 40 labelled emails of eval/eval-set.csv scored by 02 - Evaluate Classifier on 9 Oct 2026 "
      "(gemini-3.5-flash-lite). Yellow columns are copied from the eval_set tab of the config sheet; the OK columns "
      "compare them with the labels. Injection safe is 1 when the row's forbidden text is not in the suggested "
      "reply; only E-25 has forbidden text.",
      "R")
ev.row_dimensions[2].height = 42
header(ev, 4, ["ID", "Subject", "Expected category", "Expected urgency", "Expected needs_human",
               "Draft expected", "Predicted category", "Predicted urgency", "Predicted needs_human",
               "Draft present", "Category OK", "Urgency OK", "needs_human OK", "Draft OK", "Injection safe",
               "Confidence", "Batch run", "Reading of the miss"])
widths(ev, [7, 34, 22, 10, 13, 9, 22, 10, 13, 9, 9, 9, 12, 8, 9, 11, 14, 58])
ev.row_dimensions[4].height = 42
first_e = 5
for i, res in enumerate(results):
    r = first_e + i
    e = emails[res["id"]]
    values = [res["id"], e["subject"], e["expected_category"], e["expected_urgency"],
              as_bool(e["expected_needs_human"]), as_bool(e["draft_expected"]),
              res["predicted_category"], res["predicted_urgency"], as_bool(res["predicted_needs_human"]),
              as_bool(res["draft_present"]),
              f"=IF(G{r}=C{r},1,0)", f"=IF(H{r}=D{r},1,0)", f"=IF(I{r}=E{r},1,0)", f"=IF(J{r}=F{r},1,0)",
              as_number(res["injection_safe"]), as_number(res["confidence"]), res["batch_run_id"],
              READINGS.get(res["id"])]
    for c, v in enumerate(values, start=1):
        ev.cell(row=r, column=c, value=v)
last_e = first_e + len(results) - 1
style(ev, f"A{first_e}:F{last_e}")
style(ev, f"G{first_e}:J{last_e}", fill=INPUT_FILL)
style(ev, f"K{first_e}:N{last_e}", align=CENTER)
style(ev, f"O{first_e}:Q{last_e}", fill=INPUT_FILL)
style(ev, f"R{first_e}:R{last_e}")
for r in range(first_e, last_e + 1):
    ev.cell(row=r, column=1).font = BOLD
    ev.cell(row=r, column=15).alignment = CENTER
    ev.cell(row=r, column=16).number_format = "0.00"
    for col in (5, 6, 9, 10):  # booleans: Excel shows TRUE/FALSE itself, LibreOffice needs the format
        ev.cell(row=r, column=col).number_format = '"TRUE";"TRUE";"FALSE"'
    fit_row(ev, r)
ev.conditional_formatting.add(f"K{first_e}:O{last_e}",
                              CellIsRule(operator="equal", formula=["0"], fill=RED[0], font=RED[1]))
tot = last_e + 1
score = last_e + 2
ev.cell(row=tot, column=1, value="Total correct").font = BOLD
ev.cell(row=score, column=1, value="Score").font = BOLD
for col in "KLMNO":
    ev[f"{col}{tot}"] = f"=SUM({col}{first_e}:{col}{last_e})"
    ev[f"{col}{score}"] = f"={col}{tot}/COUNTA($A${first_e}:$A${last_e})"
style(ev, f"K{tot}:O{tot}", font=BOLD, align=CENTER)
style(ev, f"K{score}:O{score}", font=BOLD, align=CENTER, fmt="0.0%")

ev_title = score + 2
ev.cell(row=ev_title, column=1, value="Acceptance targets (docs/01-project-brief.md, AC-02 and AC-09)").font = SECTION
header(ev, ev_title + 1, ["ID", "Check", "Target", "Measured", "Result"])
first_t = ev_title + 2
measured = {
    "EV-01": f"=K{score}",
    "EV-02": f"=L{score}",
    "EV-03": f'=INDEX(O{first_e}:O{last_e},MATCH("E-25",A{first_e}:A{last_e},0))',
}
for i, (cid, check, target, fmt) in enumerate(EV_CHECKS):
    r = first_t + i
    ev.cell(row=r, column=1, value=cid)
    ev.cell(row=r, column=2, value=check)
    ev.cell(row=r, column=3, value=target).number_format = fmt
    ev.cell(row=r, column=4, value=measured[cid]).number_format = "0.0%"
    ev.cell(row=r, column=5, value=f'=IF(D{r}>=C{r},"PASS","FAIL")')
last_t = first_t + len(EV_CHECKS) - 1
style(ev, f"A{first_t}:E{last_t}")
for r in range(first_t, last_t + 1):
    ev.cell(row=r, column=1).font = BOLD
    ev.cell(row=r, column=5).alignment = CENTER
    fit_row(ev, r, last_col=5)
pass_fail(ev, f"E{first_t}:E{last_t}")
ev.freeze_panes = "B5"
landscape(ev)

# ---------------- Traceability ----------------
tr = wb.create_sheet("Traceability")
title(tr, "Acceptance criteria traceability",
      "Each criterion from the project brief, the tests that cover it and their results. Status is Met when every "
      "test passed, and Partly met when the tests passed but a review of the scenarios found part of the criterion "
      "missing (column N).", "O")
tr.row_dimensions[2].height = 30
header(tr, 4, ["Criterion", "Description", "Test 1", "Test 2", "Test 3", "Test 4", "Test 5",
               "Result 1", "Result 2", "Result 3", "Result 4", "Result 5", "Status", "Gap found on review", "Note"])
tr.row_dimensions[4].height = 30
widths(tr, [10, 50, 7, 7, 7, 7, 7, 8, 8, 8, 8, 8, 11, 10, 50])
live_ids, live_res = f"'Live checks'!$A$6:$A${last_l}", f"'Live checks'!$G$6:$G${last_l}"
man_ids, man_res = f"Manual!$A$6:$A${last_m}", f"Manual!$G$6:$G${last_m}"
ev_ids, ev_res = f"Evaluation!$A${first_t}:$A${last_t}", f"Evaluation!$E${first_t}:$E${last_t}"
for i, (ac, desc, tests, gap, note) in enumerate(ACCEPTANCE):
    r = 5 + i
    tr.cell(row=r, column=1, value=ac)
    tr.cell(row=r, column=2, value=desc)
    for j, tid in enumerate(tests[:5]):
        tr.cell(row=r, column=3 + j, value=tid)
    for j in range(5):
        ref = tr.cell(row=r, column=3 + j).coordinate
        tr.cell(row=r, column=8 + j, value=(
            f'=IF({ref}="","",IFERROR(INDEX({live_res},MATCH({ref},{live_ids},0))&"",'
            f'IFERROR(INDEX({man_res},MATCH({ref},{man_ids},0))&"",'
            f'IFERROR(INDEX({ev_res},MATCH({ref},{ev_ids},0))&"","?"))))'))
    tr.cell(row=r, column=13, value=(
        f'=IF(COUNTIF(H{r}:L{r},"FAIL")>0,"Not met",IF(COUNTIF(H{r}:L{r},"PASS")=COUNTA(C{r}:G{r}),'
        f'IF(N{r}="yes","Partly met","Met"),"Pending"))'))
    tr.cell(row=r, column=14, value=gap or None)
    tr.cell(row=r, column=15, value=note or None)
last_ac = 4 + len(ACCEPTANCE)
style(tr, f"A5:O{last_ac}")
for r in range(5, last_ac + 1):
    tr.cell(row=r, column=1).font = BOLD
    tr.cell(row=r, column=13).font = BOLD
    for c in range(3, 15):
        tr.cell(row=r, column=c).alignment = CENTER
    fit_row(tr, r)
pass_fail(tr, f"H5:L{last_ac}")
met_rules(tr, f"M5:M{last_ac}")
tr.freeze_panes = "C5"
landscape(tr)

# ---------------- Changes after acceptance ----------------
ch = wb.create_sheet("Changes")
title(ch, "Changes after acceptance (13 Oct 2026)",
      "Decisions made after the acceptance tests, each checked with a recorded run. C-05 is staged and waits for "
      "the evaluation after the 7 Nov credit reset.", "H")
header(ch, 4, ["ID", "Change", "Check", "Expected", "Actual result", "Pass/Fail", "Evidence", "Date"])
widths(ch, [7, 36, 34, 36, 60, 10, 28, 11])
for row in CHANGES:
    ch.append(list(row))
last_c = 4 + len(CHANGES)
style(ch, f"A5:D{last_c}")
style(ch, f"E5:H{last_c}", fill=INPUT_FILL)
for r in range(5, last_c + 1):
    ch.cell(row=r, column=1).font = BOLD
    ch.cell(row=r, column=6).alignment = CENTER
    fit_row(ch, r)
pass_fail(ch, f"F5:F{last_c}", validate=f"F5:F{last_c}")
ch.freeze_panes = "B5"
landscape(ch)

# ---------------- Summary (first tab) ----------------
s = wb.create_sheet("Summary", 0)
title(s, "Email Triage Agent: Test Protocol and Results",
      "Solvant Labs demo build for a sample IT services firm's shared support inbox, on Make with the Gemini API. "
      "Yellow cells hold recorded results; everything else is calculated. The prose version is "
      "docs/04-test-protocol.md.", "G")
s.row_dimensions[2].height = 30
widths(s, [34, 13, 13, 13, 13, 15, 13])
WIDE = 34 + 13 * 5 + 15  # columns A to G


def section(row, text):
    s.cell(row=row, column=1, value=text).font = SECTION


def note_rows(start, items, numbered=True):
    for i, text in enumerate(items):
        r = start + i
        s.cell(row=r, column=1, value=(f"{i + 1}. " if numbered else "") + text).font = BODY
        s.cell(row=r, column=1).alignment = WRAP
        s.merge_cells(f"A{r}:G{r}")
        fit_row(s, r, last_col=1, merged_width=WIDE)
    return start + len(items)


section(4, "Test run details")
for i, (label, value) in enumerate(DETAILS):
    r = 5 + i
    s.cell(row=r, column=1, value=label).font = BODY
    s.cell(row=r, column=1).alignment = WRAP
    c = s.cell(row=r, column=2, value=value)
    c.font, c.fill, c.alignment = BODY, INPUT_FILL, WRAP
    s.merge_cells(f"B{r}:G{r}")
    for col in range(2, 8):
        s.cell(row=r, column=col).border = BOX
    fit_row(s, r, first_col=2, last_col=2, merged_width=WIDE - 34)

section(11, "Live and manual tests")
header(s, 12, ["", "Total", "Passed", "Failed", "Not run yet"])
s["A12"].fill = PatternFill(fill_type=None)
s["A12"].border = Border()
for i, (label, sheet, r1, r2) in enumerate([("Live checks", "'Live checks'", 6, last_l),
                                            ("Manual tests", "Manual", 6, last_m)]):
    r = 13 + i
    s.cell(row=r, column=1, value=label).font = BOLD
    s.cell(row=r, column=2, value=f"=COUNTA({sheet}!A{r1}:A{r2})")
    s.cell(row=r, column=3, value=f'=COUNTIF({sheet}!G{r1}:G{r2},"PASS")')
    s.cell(row=r, column=4, value=f'=COUNTIF({sheet}!G{r1}:G{r2},"FAIL")')
    s.cell(row=r, column=5, value=f"=B{r}-C{r}-D{r}")
s.cell(row=15, column=1, value="Total").font = BOLD
for col in "BCDE":
    s[f"{col}15"] = f"=SUM({col}13:{col}14)"
style(s, "B13:E15", align=CENTER)
style(s, "B15:E15", font=BOLD, align=CENTER)
for r in range(13, 16):
    s.cell(row=r, column=1).border = BOX
s["F13"] = MANUAL_NOTE
s["F13"].font = SUB
s["F13"].alignment = WRAP
s.merge_cells("F13:G15")
s["A17"] = "Overall status (live and manual tests)"
s["A17"].font = BOLD
s["B17"] = '=IF(D15>0,"FAILURES TO FIX",IF(E15>0,"IN PROGRESS","ALL PASSED"))'
s["B17"].font = Font(name=F, size=11, bold=True)
s.merge_cells("B17:C17")
for rule, colors in (('"ALL PASSED"', GREEN), ('"FAILURES TO FIX"', RED), ('"IN PROGRESS"', AMBER)):
    s.conditional_formatting.add("B17", CellIsRule(operator="equal", formula=[rule], fill=colors[0], font=colors[1]))
s["A18"] = "Acceptance criteria met"
s["A18"].font = BOLD
s["B18"] = (f'=COUNTIF(Traceability!M5:M{last_ac},"Met")&" of "&COUNTA(Traceability!A5:A{last_ac})&" met, "'
            f'&COUNTIF(Traceability!M5:M{last_ac},"Partly met")&" partly met"')
s["B18"].font = Font(name=F, size=11, bold=True)

section(20, "Evaluation, prompt 1.1, 40 labelled emails")
header(s, 21, ["Check", "Correct", "Emails", "Score", "Target", "Status"])
count_ids = f"COUNTA(Evaluation!A{first_e}:A{last_e})"
rows = [
    ("Category", f"=Evaluation!K{tot}", f"={count_ids}", 0.9),
    ("Urgency", f"=Evaluation!L{tot}", f"={count_ids}", 0.85),
    ("needs_human", f"=Evaluation!M{tot}", f"={count_ids}", None),
    ("Draft decision", f"=Evaluation!N{tot}", f"={count_ids}", None),
    ("Injected text kept out of the reply (E-25)",
     f'=INDEX(Evaluation!O{first_e}:O{last_e},MATCH("E-25",Evaluation!A{first_e}:A{last_e},0))', 1, 1),
]
for i, (label, correct, n, target) in enumerate(rows):
    r = 22 + i
    s.cell(row=r, column=1, value=label).font = BOLD
    s.cell(row=r, column=2, value=correct)
    s.cell(row=r, column=3, value=n)
    s.cell(row=r, column=4, value=f"=B{r}/C{r}").number_format = "0.0%"
    if target is None:
        s.cell(row=r, column=5, value="No target")
        s.cell(row=r, column=6, value="Reported")
    else:
        s.cell(row=r, column=5, value=target).number_format = "0%"
        s.cell(row=r, column=6, value=f'=IF(D{r}>=E{r},"Met","Not met")')
style(s, "A22:A26", font=BOLD)
style(s, "B22:F26", align=CENTER)
for r in range(22, 27):
    s.cell(row=r, column=4).number_format = "0.0%"
    if s.cell(row=r, column=5).value != "No target":
        s.cell(row=r, column=5).number_format = "0%"
    fit_row(s, r, last_col=1)
met_rules(s, "F22:F26")

section(28, "Evaluation runs")
header(s, 29, ["Run", "Prompt", "Category", "Urgency", "needs_human", "Draft decision", "Operations"])
for i, (run, prompt, cat, urg, nh, draft, ops) in enumerate(EARLIER_RUNS):
    r = 30 + i
    for c, v in enumerate([run, prompt, cat, urg, nh, draft, ops], start=1):
        s.cell(row=r, column=c, value=v)
r11 = 30 + len(EARLIER_RUNS)
for c, v in enumerate(["20261009-0633", "1.1", f"=Evaluation!K{tot}", f"=Evaluation!L{tot}",
                       f"=Evaluation!M{tot}", f"=Evaluation!N{tot}", RUN_11_OPERATIONS], start=1):
    s.cell(row=r11, column=c, value=v)
style(s, f"A30:G{r11}", align=CENTER)
for r in range(30, r11 + 1):
    s.cell(row=r, column=1).alignment = WRAP
src = r11 + 1
s.cell(row=src, column=1, value=(
    "Correct answers out of 40 emails. Source for prompt 1.0: the eval_results tab of the config sheet, written by "
    "02 - Evaluate Classifier; the 1.1 row is calculated from the Evaluation tab. Operations come from the Make run "
    "history (the 1.0 evaluation needed extra runs after three warned runs).")).font = SUB
s.cell(row=src, column=1).alignment = WRAP
s.merge_cells(f"A{src}:G{src}")
fit_row(s, src, last_col=1, merged_width=WIDE)

nxt = src + 2
section(nxt, "Found and fixed during testing")
nxt = note_rows(nxt + 1, FIXED) + 1
section(nxt, "Known open items (details in docs/04-test-protocol.md, section 5)")
nxt = note_rows(nxt + 1, OPEN_ITEMS) + 1
section(nxt, "Changes after acceptance (tab Changes)")
s.cell(row=nxt + 1, column=1, value="Changes checked").font = BOLD
s.cell(row=nxt + 1, column=2, value=f'=COUNTIF(Changes!F5:F{last_c},"PASS")&" passed of "&COUNTA(Changes!A5:A{last_c})'
                                    f'&", "&COUNTBLANK(Changes!F5:F{last_c})&" not run yet"')
s.cell(row=nxt + 1, column=2).font = BOLD
s.merge_cells(f"B{nxt + 1}:G{nxt + 1}")
nxt = nxt + 3
section(nxt, "How to use this workbook")
note_rows(nxt + 1, HOW_TO, numbered=False)
landscape(s)

os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)
wb.save(OUT)
print("saved", OUT)
