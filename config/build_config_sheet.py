"""Build config/triage-config.xlsx: the Triage Config sheet the team edits (upload to Google Sheets)."""
import csv
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(HERE, 'triage-config.xlsx')

HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='101820')
EDIT_FILL = PatternFill('solid', fgColor='FFF4CC')  # yellow: the team edits these
WRAP = Alignment(wrap_text=True, vertical='top')


def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(1, i); c.font = HEAD; c.fill = HEAD_FILL; c.alignment = WRAP
    ws.freeze_panes = 'A2'


wb = openpyxl.Workbook()

# settings
ws = wb.active; ws.title = 'settings'
header(ws, ['key', 'value', 'what it does'], [28, 46, 70])
rows = [
    ('team_name', 'Solvant Labs Support', 'How replies are signed and who the prompt says you are'),
    ('slack_channel', '#support-triage', 'For reference only for now: the scenarios post to the channel set in live-ids.json when the blueprints are generated (open item 10)'),
    ('confidence_threshold', 0.7, 'Below this the email is routed to Review instead of its category'),
    ('model', 'gemini-3.5-flash-lite', 'Gemini model ID; 00 - Connection Check lists the ones your key can use'),
    ('max_reply_words', 120, 'Not used yet: the prompt itself sets the 120-word limit (open item 10)'),
    ('lead_router_webhook_url', '', 'POST endpoint for sales inquiries (the n8n lead router). Blank skips the hand-off'),
    ('inbox_address', 'you+support@gmail.com', 'Where 03 - Seed Test Emails sends test emails. 01 searches the address set in live-ids.json when the blueprints are generated, so keep the two the same'),
    ('prompt_version', '1.1', 'Bump when you change the prompt tab, so evaluation runs are comparable'),
    ('seed_ids', 'E-01,E-13,E-31', 'Rows of eval_set that 03 - Seed Test Emails sends to the inbox, comma separated'),
    ('max_output_tokens', 2048, 'Upper limit for the model answer in 01, in tokens. Leave at 2048; test L-09 lowers it to 30 to make the answer unreadable on purpose'),
]
for r in rows:
    ws.append(list(r))
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    row[1].fill = EDIT_FILL
    for c in row: c.alignment = WRAP

# reply_templates
ws = wb.create_sheet('reply_templates')
header(ws, ['category', 'tone', 'must_include', 'must_avoid', 'closing'], [20, 22, 60, 50, 30])
templates = [
    ('support_request', 'Calm and concrete', 'Acknowledge the problem in one line; say a technician will pick it up and what they will check first; ask for one useful detail if it is missing (how many people are affected, when it started)', 'Promising a fix time; blaming the sender; technical jargon', 'We will keep you posted.'),
    ('sales_inquiry', 'Warm and brief', 'Thank them for reaching out; one line confirming we work with companies like theirs; say a consultant will reply with next steps; ask about team size or timeline if not given', 'Prices, discounts, availability promises, naming other clients', 'Looking forward to talking.'),
    ('billing_question', 'Polite and factual', 'Acknowledge the specific invoice or request by number; say the billing team will check the account and reply; mention any document that will be needed', 'Confirming refunds, credits, amounts or cancellation terms', 'Thank you for your patience.'),
    ('scheduling', 'Friendly and short', 'Acknowledge the requested change; say someone will confirm the slot; offer to suggest alternatives', 'Confirming the slot as booked', 'Talk soon.'),
]
for r in templates:
    ws.append(list(r))
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for c in row: c.fill = EDIT_FILL; c.alignment = WRAP

# prompt
ws = wb.create_sheet('prompt')
header(ws, ['system_prompt'], [120])
prompt = open(os.path.join(ROOT, 'prompts', 'triage-system-prompt.md'), encoding='utf-8').read().strip()
ws.append([prompt]); ws['A2'].alignment = WRAP; ws['A2'].fill = EDIT_FILL
ws.row_dimensions[2].height = 400
ws.append(['The scenarios read this cell at the start of every run. {{team_name}} and {{reply_templates}} are filled from the other tabs. Change the text, bump prompt_version in settings, then run 02 - Evaluate Classifier to see the effect before it reaches the inbox.'])
ws['A3'].alignment = WRAP; ws['A3'].font = Font(italic=True, color='4A5560')

# eval_set
ws = wb.create_sheet('eval_set')
cols = ['id', 'from_name', 'from_email', 'subject', 'body', 'expected_category', 'expected_urgency', 'expected_needs_human', 'draft_expected', 'forbidden_in_reply', 'notes',
        'predicted_category', 'predicted_urgency', 'predicted_needs_human', 'draft_present', 'category_ok', 'urgency_ok', 'needs_human_ok', 'draft_ok', 'injection_safe', 'confidence', 'last_run_id']
header(ws, cols, [7, 16, 30, 32, 70, 20, 14, 12, 12, 22, 30, 20, 14, 12, 12, 10, 10, 12, 10, 12, 10, 12])
with open(os.path.join(ROOT, 'eval', 'eval-set.csv'), encoding='utf-8') as f:
    for r in csv.DictReader(f):
        ws.append([r[c] for c in cols[:11]])
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    row[4].alignment = WRAP
    row[10].alignment = WRAP

# triage_log
ws = wb.create_sheet('triage_log')
header(ws, ['processed_at', 'message_id', 'thread_id', 'from_email', 'from_name', 'subject', 'category', 'urgency', 'confidence', 'sentiment', 'language', 'needs_human', 'needs_human_reason', 'summary', 'key_facts', 'action', 'draft', 'slack', 'lead_router', 'model', 'input_tokens', 'output_tokens', 'raw_category', 'status', 'error', 'run_id'],
       [20, 22, 22, 30, 16, 32, 18, 10, 10, 11, 9, 11, 30, 50, 30, 18, 20, 11, 16, 20, 11, 12, 11, 10, 30, 12])

# eval_results
ws = wb.create_sheet('eval_results')
header(ws, ['run_at', 'run_id', 'model', 'prompt_version', 'emails', 'category_correct', 'category_accuracy', 'urgency_correct', 'urgency_accuracy', 'needs_human_correct', 'draft_correct', 'injection_safe', 'input_tokens', 'output_tokens', 'notes'],
       [20, 12, 22, 14, 8, 16, 16, 15, 16, 18, 13, 14, 12, 13, 40])

wb.save(OUT)
print('wrote', OUT)
